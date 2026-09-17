#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ชุดทดสอบ invoice_edenstar: ที่อยู่บริษัทต้องถูกเขียนจากค่าคงที่ + ต้องไม่มีชุดเก่าหลงเหลือ

รัน: python tools/test_invoice_edenstar_header.py
ที่มา (กับดัก 16 ก.ย. 69): สคริปต์โคลนโครงจาก template `7月前半` ที่ยังมีที่อยู่เก่า ⇒ ไฟล์ที่สร้าง
อัตโนมัติได้ที่อยู่เก่าโดยไม่มีสัญญาณเตือน · เทสต์นี้สร้าง invoice จริงลง **โฟลเดอร์ชั่วคราว**
(ไม่แตะไฟล์บน Drive) แล้วอ่านกลับมาตรวจ 5 เซลล์ + สแกนทั้งชีตหาชุดเก่า
"""
from __future__ import annotations

import datetime as dt
import os
import sys
import tempfile
from pathlib import Path

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))
import openpyxl  # noqa: E402
import invoice_edenstar as inv  # noqa: E402

fails: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        print("PASS  %s" % name)
    else:
        fails.append("%s %s" % (name, detail))
        print("FAIL  %s %s" % (name, detail))


if not os.path.exists(inv.DEFAULT_TEMPLATE):
    print("ข้าม: ไม่พบ template %s (เทสต์นี้ต้องมีไฟล์ template)" % inv.DEFAULT_TEMPLATE)
    raise SystemExit(0)

with tempfile.TemporaryDirectory(prefix="inv_hdr_test_") as tmp:
    out = os.path.join(tmp, "ทดสอบ - 請求書 - 10月前半.xlsx")
    items = [(dt.date(2026, 10, 1), "Group A"), (dt.date(2026, 10, 2), "Group B")]
    inv.build_invoice(inv.DEFAULT_TEMPLATE, out, 2026, 10, "front", items)

    ws = openpyxl.load_workbook(out).active
    for cell, want in inv.COMPANY_HEADER.items():
        got = str(ws[cell].value or "").strip()
        check("เซลล์ %s = %s" % (cell, want), got == want, "ได้ %r" % got)

    text = "\n".join(str(c.value or "") for row in ws.iter_rows() for c in row)
    for old in inv.OLD_COMPANY_VALUES:
        check("ไม่มีชุดเก่า %s" % old, old not in text)

    check("ยังมี NO (K3)", str(ws["K3"].value or "").startswith("20261015"), repr(ws["K3"].value))
    check("กำหนดชำระ 31 ต.ค. 2026",
          str(ws["D43"].value or "")[:10] == "2026-10-31", repr(ws["D43"].value))

    # verify_company_header ต้อง "จับได้" ถ้าเซลล์ถูกแก้ให้ผิด (negative control)
    wb = openpyxl.load_workbook(out)
    wb.active["I10"] = "新宿区新宿4-3-15-405"  # ใส่ที่อยู่เก่ากลับ
    bad = os.path.join(tmp, "bad.xlsx")
    wb.save(bad)
    try:
        inv.verify_company_header(bad)
        check("negative control: ตรวจจับที่อยู่ผิดได้", False, "ไม่โยน error")
    except RuntimeError as exc:
        check("negative control: ตรวจจับที่อยู่ผิดได้", True)
        print("      → %s" % str(exc)[:90])

print()
if fails:
    print("❌ ไม่ผ่าน %d ข้อ:" % len(fails))
    for f in fails:
        print("  - %s" % f)
    raise SystemExit(1)
print("✅ ผ่านทุกข้อ (test_invoice_edenstar_header)")
