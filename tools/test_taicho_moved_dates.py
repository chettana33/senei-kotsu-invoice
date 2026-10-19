"""เทสต์: ช่อง 配車日 (col N) เขียนแบบ "ย้ายวัน" เช่น '10/25→10/26' — เที่ยวต้องไปอยู่วัน **ปลายลูกศร**

ที่มา 19 ต.ค. 69: พี่เจยืนยันความหมายของ `10/25→10/26` ในใบ 1019 = ย้ายจาก 25 ไป 26
แต่ `serial_to_date()` คืน None กับข้อความ ⇒ read_request_form ข้ามแถวเงียบ ๆ (台帳 ขาดเที่ยว)
+ วันต้นลูกศรต้องถูกล้างเมื่อไม่มีแถวอื่นเหลือ (ฟอร์ม = ความจริงเต็มของวัน)

รัน: python tools/test_taicho_moved_dates.py   (ไม่ยิงเน็ต · ไม่แตะชีต)
"""
from __future__ import annotations

import datetime
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl  # noqa: E402
import taicho_gsheets as tgs  # noqa: E402

FAIL = 0


def check(label: str, got, want) -> None:
    global FAIL
    ok = got == want
    if not ok:
        FAIL += 1
    print(f"{'✅' if ok else '❌'} {label}: got={got!r} want={want!r}")


# ---- 1) arrow_dates: pure ----
check("'10/25→10/26' (ต.ค.)",
      tgs.arrow_dates("10/25\n→10/26", 10, 2026),
      (datetime.date(2026, 10, 26), datetime.date(2026, 10, 25)))
check("'10/30→10月31日'", tgs.arrow_dates("10/30\n→10月31日", 10, 2026),
      (datetime.date(2026, 10, 31), datetime.date(2026, 10, 30)))
check("'11月5日→11月6日'", tgs.arrow_dates("11月5日→\n11月6日", 11, 2026),
      (datetime.date(2026, 11, 6), datetime.date(2026, 11, 5)))
check("'1月2日→1月3日' (ม.ค. 2027)", tgs.arrow_dates("1月2日→\n1月3日", 1, 2027),
      (datetime.date(2027, 1, 3), datetime.date(2027, 1, 2)))
check("ข้ามปี '12/31→1/1' (ธ.ค. 2026)", tgs.arrow_dates("12/31→1/1", 12, 2026),
      (datetime.date(2027, 1, 1), datetime.date(2026, 12, 31)))
check("ไม่มีลูกศร = ไม่ใช่วันแบบย้าย", tgs.arrow_dates("10/25", 10, 2026), (None, None))
check("ค่าเป็น datetime = ไม่ใช่วันแบบย้าย",
      tgs.arrow_dates(datetime.datetime(2026, 10, 25), 10, 2026), (None, None))

# ---- 2) read_request_form: workbook สังเคราะห์ 3 แท็บ ----
wb = openpyxl.Workbook()
wb.remove(wb.active)

ws = wb.create_sheet("26.10")
ws["N8"] = "配車日"
rows_oct = [
    (datetime.datetime(2026, 10, 25), datetime.time(6, 0), "HTL", "AP2", "050-0001"),
    ("10/25\n→10/26", datetime.time(6, 0), "HTL", "AP2", "050-0002"),   # ย้ายวัน 25 -> 26
    (datetime.datetime(2026, 10, 26), "6:00→\n6：15", "HTL", "AP2", "050-0003"),
]
for i, (n, c, d, g, j) in enumerate(rows_oct, start=10):
    ws.cell(row=i, column=14, value=n)
    ws.cell(row=i, column=3, value=c)
    ws.cell(row=i, column=4, value=d)
    ws.cell(row=i, column=7, value=g)
    ws.cell(row=i, column=10, value=j)

ws = wb.create_sheet("26.11")
ws["N8"] = "配車日"
ws.cell(row=10, column=14, value="11月5日→\n11月6日")   # แถวเดียวของวันต้นทาง = ต้องล้าง 11/5
ws.cell(row=10, column=3, value=datetime.time(6, 45))
ws.cell(row=10, column=4, value="HTL")
ws.cell(row=10, column=7, value="AP1")

ws = wb.create_sheet("27.01")
ws["N8"] = "配車日"
ws.cell(row=10, column=14, value="1月2日→\n1月3日")
ws.cell(row=10, column=3, value="17:50→\n18：30")
ws.cell(row=10, column=4, value="HTL")
ws.cell(row=10, column=7, value="AP2")

tmp = os.path.join(tempfile.gettempdir(), "test_taicho_moved_dates.xlsx")
wb.save(tmp)
try:
    out = tgs.read_request_form(tmp)
    e = lambda t, d, g: {"t": t, "h": None, "D": d, "G": g}  # noqa: E731
    check("10/25 เหลือแถวเดียว (แถวย้ายวันไม่นับที่วันต้นทาง)",
          out.get(datetime.date(2026, 10, 25)), [e("06:00", "HTL", "AP2")])
    check("10/26 ได้ 2 เที่ยว (06:00 จากแถวย้ายวัน + 06:15)",
          out.get(datetime.date(2026, 10, 26)),
          [e("06:00", "HTL", "AP2"), e("06:15", "HTL", "AP2")])
    check("11/5 ถูกฝากเป็น [] (วันต้นทางไม่มีแถวเหลือ = ล้าง)",
          out.get(datetime.date(2026, 11, 5)), [])
    check("11/6 ได้ 06:45 จากแถว 11月5日→11月6日",
          out.get(datetime.date(2026, 11, 6)), [e("06:45", "HTL", "AP1")])
    check("1/3 ได้ 18:30 (จาก '1月2日→1月3日' + เวลาปลายลูกศร)",
          out.get(datetime.date(2027, 1, 3)), [e("18:30", "HTL", "AP2")])
    check("1/2 ถูกฝากเป็น []", out.get(datetime.date(2027, 1, 2)), [])
finally:
    os.remove(tmp)

print()
if FAIL:
    print(f"❌ ไม่ผ่าน {FAIL} ข้อ")
    raise SystemExit(1)
print("✅ ผ่านทุกข้อ")
